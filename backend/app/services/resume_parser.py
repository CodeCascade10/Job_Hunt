
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader


MAX_RESUME_SIZE = 5 * 1024 * 1024  # 5 MB
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class ResumeParsingError(ValueError):
    """Raised when a resume cannot be safely parsed."""


def extract_resume_text(filename: str, content: bytes) -> str:
    """Extract text locally from a PDF or DOCX resume."""

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeParsingError(
            "Unsupported file format. Upload a PDF or DOCX file."
        )

    if not content:
        raise ResumeParsingError("The uploaded file is empty.")

    if len(content) > MAX_RESUME_SIZE:
        raise ResumeParsingError(
            "Resume exceeds the maximum allowed size of 5 MB."
        )

    if extension == ".pdf":
        return _extract_pdf_text(content)

    return _extract_docx_text(content)


def _extract_pdf_text(content: bytes) -> str:
    if not content.startswith(b"%PDF-"):
        raise ResumeParsingError("The uploaded file is not a valid PDF.")

    try:
        reader = PdfReader(BytesIO(content), strict=True)

        if reader.is_encrypted:
            raise ResumeParsingError(
                "Password-protected PDFs are not supported."
            )

        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")

        text = "\n".join(pages).strip()

    except ResumeParsingError:
        raise
    except Exception as exc:
        raise ResumeParsingError(
            "Could not read this PDF. Check that the file is valid."
        ) from exc

    if not text:
        raise ResumeParsingError(
            "No readable text was found. This may be a scanned PDF; "
            "OCR support is not implemented yet."
        )

    return text


def _extract_docx_text(content: bytes) -> str:
    try:
        with ZipFile(BytesIO(content)) as archive:
            if archive.testzip() is not None:
                raise ResumeParsingError("The DOCX file is corrupted.")

            if "word/document.xml" not in archive.namelist():
                raise ResumeParsingError("The DOCX document is invalid.")

        document = Document(BytesIO(content))

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Include text inside tables, which many resumes use for layout.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())

        text = "\n".join(parts).strip()

    except ResumeParsingError:
        raise
    except (BadZipFile, Exception) as exc:
        raise ResumeParsingError(
            "Could not read this DOCX. Check that the file is valid."
        ) from exc

    if not text:
        raise ResumeParsingError(
            "No readable text was found in the DOCX document."
        )

    return text
