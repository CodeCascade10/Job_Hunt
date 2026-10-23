
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from app.services.resume_parser import (
    MAX_RESUME_SIZE,
    ResumeParsingError,
    extract_resume_text,
)


def make_docx_bytes() -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Alex Morgan")
    document.add_paragraph("Python Developer")
    document.add_paragraph("Skills: Python, FastAPI")
    document.save(buffer)
    return buffer.getvalue()


def make_blank_pdf_bytes() -> bytes:
    buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(buffer)
    return buffer.getvalue()


def test_extract_text_from_docx():
    content = make_docx_bytes()

    text = extract_resume_text("resume.docx", content)

    assert "Alex Morgan" in text
    assert "Python Developer" in text
    assert "Python, FastAPI" in text


def test_reject_unsupported_file_extension():
    with pytest.raises(ResumeParsingError, match="Unsupported file format"):
        extract_resume_text("resume.txt", b"some text")


def test_reject_empty_file():
    with pytest.raises(ResumeParsingError, match="empty"):
        extract_resume_text("resume.pdf", b"")


def test_reject_file_over_size_limit():
    content = b"x" * (MAX_RESUME_SIZE + 1)

    with pytest.raises(ResumeParsingError, match="maximum allowed size"):
        extract_resume_text("resume.docx", content)


def test_reject_invalid_pdf_signature():
    with pytest.raises(ResumeParsingError, match="not a valid PDF"):
        extract_resume_text("resume.pdf", b"not a PDF file")


def test_reject_invalid_docx():
    with pytest.raises(ResumeParsingError, match="Could not read this DOCX"):
        extract_resume_text("resume.docx", b"not a DOCX file")


def test_blank_pdf_reports_missing_readable_text():
    content = make_blank_pdf_bytes()

    with pytest.raises(ResumeParsingError, match="No readable text"):
        extract_resume_text("resume.pdf", content)
